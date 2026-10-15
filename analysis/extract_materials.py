from pathlib import Path
import json
import zipfile
from datetime import date, datetime
from collections import Counter
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageOps, ImageDraw

root = Path(r'D:\DemoExam\КИМ_ДЭ_09.02.07-2')
out = Path(r'D:\DemoExam\analysis')
out.mkdir(exist_ok=True)
inventory = []
for path in sorted(root.rglob('*')):
    if not path.is_file():
        continue
    entry = {'path': str(path), 'bytes': path.stat().st_size}
    if path.suffix == '.pdf':
        reader = PdfReader(path)
        pages = [{'page': i+1, 'text': p.extract_text() or ''} for i,p in enumerate(reader.pages)]
        (out/'main_pdf.txt').write_text('\n\n'.join(f'=== PAGE {p["page"]} ===\n{p["text"]}' for p in pages), encoding='utf-8')
        (out/'main_pdf.json').write_text(json.dumps(pages, ensure_ascii=False, indent=2), encoding='utf-8')
        entry['pages'] = len(pages)
    elif path.suffix == '.docx':
        doc = Document(path)
        lines = []
        for child in doc.element.body:
            tag = child.tag.rsplit('}',1)[-1]
            if tag == 'p':
                text = ''.join(child.itertext())
                # python-docx exposes paragraph text without repeated XML run text.
                from docx.text.paragraph import Paragraph
                text = Paragraph(child,doc).text
                if text: lines.append(text)
            elif tag == 'tbl':
                from docx.table import Table
                lines.extend(' | '.join(c.text for c in r.cells) for r in Table(child,doc).rows)
        for section in doc.sections:
            for hf in [section.header, section.footer]:
                lines.extend(p.text for p in hf.paragraphs if p.text)
        (out/(path.stem+'.txt')).write_text('\n'.join(lines), encoding='utf-8')
        media_dir = out/path.stem
        with zipfile.ZipFile(path) as z:
            media = [n for n in z.namelist() if n.startswith('word/media/')]
            if media:
                media_dir.mkdir(exist_ok=True)
                for n in media: (media_dir/Path(n).name).write_bytes(z.read(n))
            entry['embedded_media'] = media
            entry['comments_parts'] = [n for n in z.namelist() if 'comment' in n.lower()]
    elif path.suffix == '.xlsx':
        book = load_workbook(path, data_only=False)
        sheets = []
        for sh in book:
            rows = []
            for row in sh:
                if any(c.value is not None for c in row):
                    rows.append({'row':row[0].row, 'values':[c.value.isoformat() if isinstance(c.value,(date,datetime)) else c.value for c in row]})
            formulas = [{'cell': c.coordinate,'formula': c.value} for row in sh for c in row if c.data_type=='f']
            sheets.append({'name':sh.title,'range':sh.calculate_dimension(),'rows':rows,'formulas':formulas,'merged':list(map(str,sh.merged_cells.ranges))})
        (out/(path.stem+'.json')).write_text(json.dumps(sheets,ensure_ascii=False,indent=2),encoding='utf-8')
        entry['sheets'] = [{'name': s['name'],'range':s['range'],'nonempty_rows':len(s['rows']),'formulas':len(s['formulas']),'header':s['rows'][0]['values']} for s in sheets]
    elif path.suffix in ['.png','.ico']:
        with Image.open(path) as im:
            im.load()
            entry.update({'size':im.size,'mode':im.mode})
    inventory.append(entry)
(out/'inventory.json').write_text(json.dumps(inventory,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps([e for e in inventory if not ('Images' in e['path'])], ensure_ascii=False,indent=2))
imgs = list(sorted(root.rglob('Images/*.png')))
contact = Image.new('RGB',(1000,((len(imgs)+4)//5)*150),'white')
draw = ImageDraw.Draw(contact)
for i,p in enumerate(imgs):
    with Image.open(p) as im:
        thumb = ImageOps.contain(im.convert('RGB'),(185,115))
        x,y = (i%5)*200,(i//5)*150
        contact.paste(thumb,(x+(200-thumb.width)//2,y))
        draw.text((x+5,y+118),p.stem,fill='black')
contact.save(out/'images_contact_sheet.png')
