from pathlib import Path
import re

root = Path(r'D:\DemoExam')
path = root / 'EXAM_LINUX_GUIDE.md'
text = path.read_text(encoding='utf-8')

intro = '''# Выполнение экзамена Чудо Обувь с нуля

Python, Tkinter + ttk, PostgreSQL и pgAdmin. Инструкция обновлена 13 октября 2026 года для школьной «М ОС» Linux.

Предполагается, что на экзамене тебе дают исходные материалы задания, а PostgreSQL, pgAdmin и Python уже установлены. Свой готовый проект или дополнительные скрипты приносить не требуется. Все файлы проекта ты создаёшь сам по ходу выполнения. Названия файлов в шагах ниже — это файлы, которые надо написать, а не скачать из моего комплекта.

По вашей договорённости Автотеста не будет, проверка предположительно ручная. Функции, порядок модулей и обязательные результаты берём из КИМ; конкретное место передачи файлов определяют преподаватели.

## 0. Создать проект и установить библиотеки

Открой терминал в папке, где собираешься работать. Создай папку проекта:

```bash
mkdir chudo_obuv
cd chudo_obuv
python3 -m venv .venv
source .venv/bin/activate
python -m pip install "psycopg[binary]" openpyxl Pillow
```

Это вся установка для приложения и чтения Excel. Psycopg подключает PostgreSQL, openpyxl читает XLSX, Pillow работает с изображениями. Tkinter и ttk используются через стандартную библиотеку, отдельно через pip их не ставим.

Создай в редакторе обычный `requirements.txt` и впиши:

```text
psycopg[binary]
openpyxl
Pillow
```

Этот файл нужен, чтобы проверяющий мог установить те же библиотеки командой `python -m pip install -r requirements.txt`. Для твоей первоначальной установки он не обязателен: библиотеки уже установлены командой выше. Для модуля 5, когда понадобятся аналитические CSV, достаточно отдельно выполнить `python -m pip install pandas`.

В IDE выбери интерпретатор `.venv/bin/python`. После открытия нового терминала снова выполни `source .venv/bin/activate`.

Если нужно проверить Tkinter перед началом работы, запусти `python -m tkinter`: должно появиться маленькое окно. Если Tkinter отсутствует, это надо исправить на школьном компьютере с преподавателем в день подготовки.

Создай через файловый менеджер папки `app`, `data`, `assets` и `database` внутри проекта. В `data` скопируй **пять Excel, выданных в задании**. В `assets/products` скопируй выданные фотографии обуви. Логотип положи в `assets/logo.png`, заглушку — в `assets/picture.png`. Копируется только то, что выдали организаторы; предварительно написанные скрипты не нужны.

На Linux учитывай три вещи: имена файлов чувствительны к регистру; пути вычисляй через pathlib относительно проекта; Windows-venv не переноси — окружение создаётся на Linux. Шрифт по ТЗ — Calibri: его доступность лучше выяснить в день подготовки. Если его нет, замену согласовать с преподавателем.

'''
text = intro + text[text.index('## 1. Модуль 1'):]

schema = (root / 'exam_support/database/schema.sql').read_text(encoding='utf-8').strip()
schema = '\n'.join(line for line in schema.splitlines() if not line.startswith('--'))
schema_section = f'''### 1.4. Написать SQL структуры базы

Открой Query Tool у базы chudo_obuv. Напиши структуру из семи таблиц. Ниже полный пример для выбранной схемы; на тренировке разберись, что делает каждое поле и ограничение, чтобы на экзамене написать его самостоятельно.

```sql
{schema}
```

Сохрани этот текст как `database/schema.sql` кнопкой сохранения Query Tool. Выполни весь SQL кнопкой Execute. Если выделен только фрагмент, pgAdmin может выполнить только его: убери случайное выделение перед запуском всего скрипта.

В дереве Schemas → public → Tables нажми Refresh. Должно появиться семь таблиц. Если скрипт дал ошибку внутри транзакции, прочитай первую ошибку, выполни `ROLLBACK;`, исправь причину. Для повторной тренировки используй отдельную новую БД, чтобы не удалять свою работу.

'''
start = text.index('### 1.4.')
end = text.index('### 1.5.', start)
text = text[:start] + schema_section + text[end:]

import_code = '''from pathlib import Path
from datetime import datetime
from decimal import Decimal

from openpyxl import load_workbook
from app.db import get_connection

DATA_DIR = Path(__file__).resolve().parent / "data"


def clean(value):
    return " ".join(str(value).replace("\\u00a0", " ").split())


def number(value):
    return Decimal(str(value).replace(",", "."))


def model_key(name, manufacturer):
    name = clean(name)
    manufacturer = clean(manufacturer)
    if name == "Черные туфли в классическом стиле" and manufacturer == "Барбари":
        name = "Черные туфли в классическом стиле — база для деловых образов"
    return name, manufacturer


def read_excel(filename):
    book = load_workbook(DATA_DIR / filename, read_only=True, data_only=True)
    try:
        return [row for row in book.worksheets[0].iter_rows(min_row=2, values_only=True)
                if any(value is not None for value in row)]
    finally:
        book.close()


def insert_id(cursor, query, values):
    cursor.execute(query, values)
    return cursor.fetchone()[0]


def import_data():
    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute("SELECT COUNT(*) FROM roles")
            if cursor.fetchone()[0] != 0:
                raise ValueError("Импорт рассчитан на новую пустую базу")

            roles = {}
            for code, name in [("admin", "Администратор"),
                               ("manager", "Менеджер"),
                               ("customer", "Авторизованный пользователь")]:
                roles[name] = insert_id(cursor,
                    "INSERT INTO roles (role_code, role_name) VALUES (%s, %s) RETURNING role_id",
                    (code, name))

            products = {}
            for category, subcategory, image, name, manufacturer, description, composition, price in read_excel("Products_import.xlsx"):
                key = model_key(name, manufacturer)
                if key in products:
                    raise ValueError(f"Повтор модели: {key}")
                products[key] = insert_id(cursor,
                    """INSERT INTO products
                       (category, subcategory, image_path, product_name, manufacturer,
                        description, composition, price)
                       VALUES (%s, %s, %s, %s, %s, %s, %s, %s) RETURNING product_id""",
                    (clean(category), clean(subcategory), image, key[0], key[1],
                     description or "", composition or "", number(price)))

            sizes = {}
            for row in read_excel("Sizes_import.xlsx"):
                value = number(row[0])
                sizes[value] = insert_id(cursor,
                    "INSERT INTO sizes (size_value) VALUES (%s) RETURNING size_id", (value,))

            users = {}
            for last_name, first_name, middle_name, login, role in read_excel("Users_import.xlsx"):
                fio = " ".join(clean(value) for value in (last_name, first_name, middle_name))
                if fio in users:
                    raise ValueError(f"Неоднозначное ФИО: {fio}")
                users[fio] = insert_id(cursor,
                    """INSERT INTO users (last_name, first_name, middle_name, login, role_id)
                       VALUES (%s, %s, %s, %s, %s) RETURNING user_id""",
                    (clean(last_name), clean(first_name), clean(middle_name), clean(login), roles[clean(role)]))

            stock_items = {}
            for name, manufacturer, size, quantity in read_excel("Stock_Items_import.xlsx"):
                key = (products[model_key(name, manufacturer)], sizes[number(size)])
                stock_items[key] = insert_id(cursor,
                    """INSERT INTO stock_items (product_id, size_id, quantity_available)
                       VALUES (%s, %s, %s) RETURNING stock_item_id""", (*key, int(quantity)))

            orders = {}
            for order_id, order_date, fio, category, name, manufacturer, size, quantity, price in read_excel("Orders_import.xlsx"):
                order_id = int(order_id)
                if isinstance(order_date, datetime):
                    order_date = order_date.date()
                user_id = users[clean(fio)]
                if order_id not in orders:
                    cursor.execute("INSERT INTO orders (order_id, order_date, user_id) VALUES (%s, %s, %s)",
                                   (order_id, order_date, user_id))
                    orders[order_id] = (order_date, user_id)
                elif orders[order_id] != (order_date, user_id):
                    raise ValueError(f"Разные клиент/дата у заказа {order_id}")
                key = (products[model_key(name, manufacturer)], sizes[number(size)])
                cursor.execute("""INSERT INTO order_items (order_id, stock_item_id, quantity, unit_price)
                                  VALUES (%s, %s, %s, %s)""",
                               (order_id, stock_items[key], int(quantity), number(price)))

            cursor.execute("""SELECT setval(pg_get_serial_sequence('public.orders', 'order_id'),
                              COALESCE(MAX(order_id), 1), COUNT(*) > 0) FROM orders""")
    print("Импорт завершён")


if __name__ == "__main__":
    import_data()
'''
db_code = '''import psycopg


def get_connection():
    return psycopg.connect(
        host="127.0.0.1",
        port=5432,
        dbname="chudo_obuv",
        user="exam_user",
        password="ПАРОЛЬ_POSTGRESQL",
        connect_timeout=5,
    )
'''
import_section = f'''### 1.6. Написать подключение и импорт Excel

Создай пустой файл `app/__init__.py`. Он обозначает папку app как Python-пакет. Затем создай `app/db.py` и напиши подключение:

```python
{db_code.rstrip()}
```

Замени имя пользователя и пароль на выданные параметры PostgreSQL. `exam_user` — пример, создавать отдельного пользователя ради этого имени не нужно. Если в pgAdmin используешь postgres или другую выданную учётную запись, укажи её. Это пароль подключения к БД, а не пароль пользователя магазина.

В корне проекта создай `import_data.py`. Ниже пример импорта для **конкретных пяти Excel этого задания**. Он читает исходники из data, получает ID вставленных строк через RETURNING и строит словари для внешних ключей.

```python
{import_code.rstrip()}
```

Это код, который ты пишешь в процессе работы, а не файл, который предполагается принести заранее. В тренировке сначала разбери порядок: роли → модели → размеры → пользователи → остатки → шапки заказов и их позиции. Не требуется выучить каждую строку; требуется понимать INSERT, RETURNING, словари сопоставления и связи.

`clean` убирает неразрывный/лишний пробел. `model_key` учитывает сокращённое название чёрных туфель. Размеры и цены преобразуются в Decimal. Для исходного заказа с одним номером шапка вставляется один раз. Неизвестная ссылка останавливает импорт; вся транзакция откатывается.

В настоящем приложении ввод количества будет проверяться отдельно; int(quantity) здесь используется для уже проверенных целочисленных колонок этого исходного набора.

Исходные остатки уже называются доступным количеством: исторические заказы при импорте повторно не списываем. Номер заказа импортируется явно, поэтому в конце настроена его последовательность. Остальные ID автоматически создаёт PostgreSQL.

### 1.7. Запустить импорт и проверить результат в pgAdmin

В активном venv, находясь в корне проекта, выполни:

```bash
python import_data.py
```

Если ошибка относится к файлам, проверь пять точных имён в data. Если к подключению — параметры в app/db.py. При KeyError прочитай неизвестный ключ и сравни его с исходными моделями/размерами/ФИО. Не удаляй ошибочную строку заказа ради завершения импорта.

Импорт запускается один раз в новой пустой БД. Повторный запуск в заполненной базе должен остановиться.

В Query Tool напиши проверки, например:

```sql
SELECT 'roles' AS table_name, COUNT(*) AS rows_count FROM roles
UNION ALL SELECT 'users', COUNT(*) FROM users
UNION ALL SELECT 'products', COUNT(*) FROM products
UNION ALL SELECT 'sizes', COUNT(*) FROM sizes
UNION ALL SELECT 'stock_items', COUNT(*) FROM stock_items
UNION ALL SELECT 'orders', COUNT(*) FROM orders
UNION ALL SELECT 'order_items', COUNT(*) FROM order_items;

SELECT SUM(quantity_available) FROM stock_items;
SELECT SUM(quantity) FROM order_items;

SELECT order_id, SUM(quantity * unit_price) AS total
FROM order_items
GROUP BY order_id
ORDER BY order_id;
```

Ожидается: 3 роли, 20 пользователей, 31 модель, 35 размеров, 92 товарные позиции, 10 заказов, 30 строк состава. Доступных пар — 1274, пар в старых заказах — 36. Заказ 1: сумма 4980; заказ 4: сумма 50992.

'''
start = text.index('### 1.6.')
end = text.index('### 1.8.', start)
text = text[:start] + import_section + text[end:]
text = text.replace('`schema.sql` — только структура; `seed.sql` — только начальные данные. Их оба сохраняй. Дополнительно нужен полный актуальный дамп БД, особенно после следующих модулей.',
                    '`schema.sql` — структура, которую ты сохранил из Query Tool; `import_data.py` — написанный тобой импорт. После загрузки данных сохрани полный дамп БД. Он должен отражать актуальные данные, особенно после следующих модулей.')

text = text.replace('  main.py\n  app/', '  main.py\n  import_data.py\n  data/\n    Products_import.xlsx\n    Sizes_import.xlsx\n    Stock_Items_import.xlsx\n    Users_import.xlsx\n    Orders_import.xlsx\n  app/')
text = text.replace('    seed.sql\n    check_import.sql\n', '')
text = text.replace('  config.example.json\n  config.local.json\n', '')
text = text.replace('Названия модулей — выбранный простой вариант, не обязательный шаблон КИМ.',
                    'Создавай файлы по мере появления соответствующего функционала. Названия модулей — выбранный простой вариант, не обязательный шаблон КИМ.')

start = text.index('### 2.2.')
end = text.index('### 2.3.', start)
connection_section = '''### 2.2. Использовать уже написанное подключение

Подключение находится в app/db.py, который ты создал для импорта. В остальных частях проекта получай его так:

```python
from app.db import get_connection

with get_connection() as connection:
    with connection.cursor() as cursor:
        cursor.execute("SELECT product_id, product_name FROM products")
        rows = cursor.fetchall()
```

Сначала выведи rows в терминал. Если видишь модели, Python получает данные из PostgreSQL; затем используй эти строки в интерфейсе. Отдельный проверочный файл не требуется.

Параметры пользователя передавай отдельно: `cursor.execute("... WHERE login = %s", (login,))`. Не вставляй логин/поиск в SQL через f-string. Контекст соединения Psycopg 3 сохраняет успешную транзакцию, откатывает при исключении и закрывает соединение.

'''
text = text[:start] + connection_section + text[end:]

text = text.replace('cp config.example.json config.local.json\n# В редакторе заполнить выданные параметры PostgreSQL.',
                    '# В app/db.py указать параметры выданной базы PostgreSQL.')
text = text.replace('База перед этим создаётся и восстанавливается по README. Укажи и offline-вариант установки, если wheelhouse передаётся разрешённым способом. Не писать «запустить как обычно»: проверяющий может видеть проект впервые.',
                    'База перед этим создаётся и восстанавливается по README. Укажи, где в app/db.py меняются параметры подключения. Перед передачей проекта замени реальный пароль в сдаваемом коде на понятный placeholder, например ПАРОЛЬ_POSTGRESQL; проверяющий подставляет параметры своей БД. Не писать «запустить как обычно»: проверяющий может видеть проект впервые.')
text = text.replace('Исключи из Git `.venv/`, `__pycache__/`, `*.pyc`, config.local.json, ZIP/RAR и offline_wheels. Оставь config.example.json и подробные параметры подготовки. Исходники, ресурсы, требования, ERD и полный SQL должны попасть в коммит.',
                    'В .gitignore внеси `.venv/`, `__pycache__/`, `*.pyc` и ZIP/RAR. Исходники, ресурсы, requirements.txt, ERD и полный SQL должны попасть в коммит. В app/db.py сдавай placeholder вместо реального пароля и поясни его настройку в README.')
text = text.replace('git add app assets database main.py requirements.txt config.example.json README.md .gitignore',
                    'git add app assets database main.py import_data.py requirements.txt README.md .gitignore')
text = text.replace('python -m zipfile -c ../submissions/Module_2_Catalog.zip main.py app assets database requirements.txt config.example.json README.md',
                    'python -m zipfile -c ../Module_2_Catalog.zip main.py app assets database import_data.py requirements.txt README.md')
text = text.replace('python -m zipfile -l ../submissions/Module_2_Catalog.zip',
                    'python -m zipfile -l ../Module_2_Catalog.zip')
text = text.replace('cd "$HOME/DemoExam/chudo_obuv"', 'cd /путь/к/chudo_obuv')
text = text.replace('В ZIP включать необходимые исходники/ресурсы/конфигурационный пример; `.venv` и реальный пароль не нужны. README должен объяснять, как создать config.local.json. Запуск через эти инструкции проверяй, а не считай очевидным.',
                    'В ZIP включать необходимые исходники, ресурсы, SQL и README. .venv не включай. В app/db.py вместо реального пароля оставь placeholder; README объясняет, как задать подключение. Для полной тренировки импорта можно включить и выданные Excel из data, но после восстановления SQL приложение само их не использует.')
text = text.replace('Подключи проверенный интерпретатор/создай venv, скопируй и заполни config, запусти main.py.',
                    'Создай venv, установи библиотеки, укажи подключение в app/db.py и запусти main.py.')
text = text.replace('В репозиторий отправляются обычные файлы, **не эти ZIP**. Архивы стоят в отдельной submissions.',
                    'В репозиторий отправляются обычные файлы, **не эти ZIP**. В примере ZIP сохраняется на уровень выше проекта.')
text = text.replace('Актуальные README/requirements/config example/ERD/full SQL/исходники;',
                    'Актуальные README/requirements.txt/ERD/full SQL/исходники;')
text = text.replace('## 9. Что было проверено при написании этой инструкции', '## 9. Как пользоваться инструкцией в тренировке')
start = text.index('## 9.')
text = text[:start] + '''## 9. Как пользоваться инструкцией в тренировке

Выполняй шаги с новой папкой и новой базой, создавая файлы самостоятельно. Сначала добейся импорта, затем работающего каталога, потом оформления заказа и управления заказами. После каждого модуля есть список проверок и точное имя результата.

Эта инструкция не является готовым приложением: интерфейс и бизнес-операции ты пишешь по приведённым шагам. Полный SQL структуры и пример импорта включены в сам текст, поэтому зависимость от заранее подготовленных файлов отсутствует. Разбирай код в тренировке, чтобы на экзамене воспроизвести необходимые действия.

Данные исходного комплекта ранее проверены: известны количество строк, особенности названий и связи. SQL и приведённый прямой импорт на реальном PostgreSQL здесь не запускались; доступ к школьной «М ОС» отсутствует. При тренировке проверь фактический запуск в своей БД. Основные команды и API остаются теми же, что в предыдущем разборе.
'''

with path.open('w', encoding='utf-8', newline='\n') as stream:
    stream.write(text)

# Вспомогательные материалы сохраняются только как отдельные учебные примеры.
support = root / 'exam_support'
(support / 'requirements.txt').write_text('psycopg[binary]\nopenpyxl\nPillow\n', encoding='utf-8')
readme = '''# Учебные примеры к разбору Чудо Обувь

Эта папка содержит ранее подготовленные вспомогательные SQL и скрипты. Полного приложения с окнами и заказами здесь нет.

Основная инструкция EXAM_LINUX_GUIDE.md переписана для выполнения с нуля: она не требует копировать эту папку и пользоваться этими файлами на экзамене. Для тренировки их можно читать как примеры.

Библиотеки устанавливаются обычным способом:

```bash
python3 -m venv .venv
source .venv/bin/activate
python -m pip install "psycopg[binary]" openpyxl Pillow
```

requirements.txt содержит тот же обычный список библиотек. Дополнительные check_environment.py/check_connection.py — необязательные старые диагностические примеры. prepare_import.py и database/seed.sql относятся к старому способу подготовки SQL; в новой инструкции показан самостоятельный прямой импорт из выданных Excel.

Исходный генератор SQL ранее проверен на всех пяти Excel: 3 роли, 20 пользователей, 31 модель, 35 размеров, 92 товарные позиции, 10 заказов, 30 позиций. SQL на реальном PostgreSQL и запуск на школьном Linux здесь не проверялись.
'''
(support / 'README.md').write_text(readme, encoding='utf-8')

print(f'Updated guide: {path}')
print(f'Characters: {len(text)}')
