from pathlib import Path
from datetime import datetime
from decimal import Decimal

from openpyxl import load_workbook
from app.db import get_connection
import psycopg

DATA_DIR = Path(__file__).resolve().parent / "data"


def clean(value):
    return " ".join(str(value).replace("\u00a0", " ").split())


def number(value):
    return Decimal(str(value).replace(",", "."))


def whole_number(value):
    number_value = number(value)
    if number_value != number_value.to_integral_value() or number_value < 0:
        raise ValueError(f"Ожидалось целое неотрицательное число: {value}")
    return int(number_value)


def model_key(name, manufacturer):
    name = clean(name)
    manufacturer = clean(manufacturer)
    if name == "Черные туфли в классическом стиле" and manufacturer == "Барбари":
        name = "Черные туфли в классическом стиле — база для деловых образов"
    return name, manufacturer


def read_excel(filename):
    book = load_workbook(DATA_DIR / filename, read_only=True, data_only=True)
    try:
        return [row for row in book.worksheets[0].iter_rows(min_row=2, values_only=True)
                if any(value is not None for value in row)]
    finally:
        book.close()


def insert_id(cursor, query, values):
    cursor.execute(query, values)
    return next(iter(cursor.fetchone().values()))


def import_data():
    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute("SELECT COUNT(*) AS total FROM roles")
            if cursor.fetchone()["total"] != 0:
                raise ValueError("Импорт рассчитан на новую пустую базу")

            roles = {}
            for code, name in [("admin", "Администратор"),
                               ("manager", "Менеджер"),
                               ("customer", "Авторизованный пользователь")]:
                roles[name] = insert_id(cursor,
                    "INSERT INTO roles (role_code, role_name) VALUES (%s, %s) RETURNING role_id",
                    (code, name))

            products = {}
            for category, subcategory, image, name, manufacturer, description, composition, price in read_excel("Products_import.xlsx"):
                key = model_key(name, manufacturer)
                if key in products:
                    raise ValueError(f"Повтор модели: {key}")
                products[key] = insert_id(cursor,
                    """INSERT INTO products
                       (category, subcategory, image_path, product_name, manufacturer,
                        description, composition, price)
                       VALUES (%s, %s, %s, %s, %s, %s, %s, %s) RETURNING product_id""",
                    (clean(category), clean(subcategory), image, key[0], key[1],
                     description or "", composition or "", number(price)))

            sizes = {}
            for row in read_excel("Sizes_import.xlsx"):
                value = number(row[0])
                sizes[value] = insert_id(cursor,
                    "INSERT INTO sizes (size_value) VALUES (%s) RETURNING size_id", (value,))

            users = {}
            for last_name, first_name, middle_name, login, role in read_excel("Users_import.xlsx"):
                fio = " ".join(clean(value) for value in (last_name, first_name, middle_name))
                if fio in users:
                    raise ValueError(f"Неоднозначное ФИО: {fio}")
                users[fio] = insert_id(cursor,
                    """INSERT INTO users (last_name, first_name, middle_name, login, role_id)
                       VALUES (%s, %s, %s, %s, %s) RETURNING user_id""",
                    (clean(last_name), clean(first_name), clean(middle_name), clean(login), roles[clean(role)]))

            stock_items = {}
            for name, manufacturer, size, quantity in read_excel("Stock_Items_import.xlsx"):
                key = (products[model_key(name, manufacturer)], sizes[number(size)])
                stock_items[key] = insert_id(cursor,
                    """INSERT INTO stock_items (product_id, size_id, quantity_available)
                       VALUES (%s, %s, %s) RETURNING stock_item_id""", (*key, whole_number(quantity)))

            orders = {}
            for order_id, order_date, fio, category, name, manufacturer, size, quantity, price in read_excel("Orders_import.xlsx"):
                order_id = int(order_id)
                if isinstance(order_date, datetime):
                    order_date = order_date.date()
                user_id = users[clean(fio)]
                if order_id not in orders:
                    cursor.execute("INSERT INTO orders (order_id, order_date, user_id) VALUES (%s, %s, %s)",
                                   (order_id, order_date, user_id))
                    orders[order_id] = (order_date, user_id)
                elif orders[order_id] != (order_date, user_id):
                    raise ValueError(f"Разные клиент/дата у заказа {order_id}")
                key = (products[model_key(name, manufacturer)], sizes[number(size)])
                cursor.execute("""INSERT INTO order_items (order_id, stock_item_id, quantity, unit_price)
                                  VALUES (%s, %s, %s, %s)""",
                               (order_id, stock_items[key], whole_number(quantity), number(price)))

            cursor.execute("""SELECT setval(pg_get_serial_sequence('public.orders', 'order_id'),
                              COALESCE(MAX(order_id), 1), COUNT(*) > 0) FROM orders""")
    print("Импорт завершён")


if __name__ == "__main__":
    try:
        import_data()
    except (ValueError, KeyError, OSError, psycopg.Error) as error:
        raise SystemExit(f"Импорт не выполнен: {error}") from error
