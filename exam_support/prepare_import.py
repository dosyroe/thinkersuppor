"""Подготовка проверяемого UTF-8 SQL для пяти XLSX из комплекта Чудо Обувь.

Подключение к PostgreSQL не нужно: schema.sql и полученный seed.sql
выполняются в pgAdmin в указанном порядке. Источники не изменяются.
Весь SQL создаётся только после успешной проверки ссылок.
"""

from __future__ import annotations

import argparse
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


ROLE_MAP = {
    "Администратор": (1, "admin"),
    "Менеджер": (2, "manager"),
    "Авторизованный пользователь": (3, "customer"),
}
MODEL_ALIASES = {
    ("Черные туфли в классическом стиле", "Барбари"):
        ("Черные туфли в классическом стиле — база для деловых образов", "Барбари"),
}


def normalize(value):
    if value is None:
        raise ValueError("Обязательное текстовое значение отсутствует")
    return " ".join(str(value).replace("\u00a0", " ").split())


def model_key(name, manufacturer):
    key = (normalize(name), normalize(manufacturer))
    return MODEL_ALIASES.get(key, key)


def decimal_value(value):
    try:
        result = Decimal(str(value).strip().replace(",", "."))
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f"Не число: {value!r}") from error
    if not result.is_finite():
        raise ValueError(f"Неконечное число: {value!r}")
    return result


def integer_value(value, minimum=0):
    number = decimal_value(value)
    if number != number.to_integral_value() or number < minimum:
        raise ValueError(f"Нужно целое число >= {minimum}: {value!r}")
    return int(number)


def date_value(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Не удалось распознать дату: {value!r}")


def read_rows(source, filename, column_count):
    matches = list(source.rglob(filename))
    if len(matches) != 1:
        raise ValueError(f"Ожидался ровно один {filename}, найдено {len(matches)}")
    book = load_workbook(matches[0], read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        result = []
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            if row_number == 1 or not any(value is not None for value in row):
                continue
            values = tuple(row[:column_count])
            if len(values) != column_count:
                raise ValueError(f"{filename}, строка {row_number}: не хватает столбцов")
            result.append((row_number, values))
        return result
    finally:
        book.close()


def sql_literal(value):
    if value is None:
        return "NULL"
    if isinstance(value, date):
        return f"DATE '{value.isoformat()}'"
    if isinstance(value, (int, Decimal)):
        return str(value)
    text = str(value)
    if "\x00" in text:
        raise ValueError("В строке обнаружен запрещённый нулевой символ")
    return "'" + text.replace("'", "''") + "'"


def insert_statement(table, columns, rows):
    values = ["(" + ", ".join(sql_literal(value) for value in row) + ")" for row in rows]
    return f"INSERT INTO public.{table} ({', '.join(columns)}) VALUES\n" + ",\n".join(values) + ";\n"


def build_seed(source):
    products = []
    product_ids = {}
    for number, row in read_rows(source, "Products_import.xlsx", 8):
        category, subcategory, image, name, manufacturer, description, composition, price = row
        key = model_key(name, manufacturer)
        if key in product_ids:
            raise ValueError(f"Products, строка {number}: дубликат модели {key}")
        product_id = len(products) + 1
        product_ids[key] = product_id
        products.append((product_id, normalize(category), normalize(subcategory),
                         str(image).strip() if image else None, key[0], key[1],
                         str(description or "").strip(), str(composition or "").strip(), decimal_value(price)))

    sizes = []
    size_ids = {}
    for number, row in read_rows(source, "Sizes_import.xlsx", 1):
        value = decimal_value(row[0])
        if value <= 0 or value != value.quantize(Decimal("0.1")) or value in size_ids:
            raise ValueError(f"Sizes, строка {number}: неверный/повторный размер {value}")
        size_id = len(sizes) + 1
        size_ids[value] = size_id
        sizes.append((size_id, value))

    users = []
    user_ids = {}
    logins = set()
    for number, row in read_rows(source, "Users_import.xlsx", 5):
        last_name, first_name, middle_name, login, role_name = map(normalize, row)
        fio = " ".join((last_name, first_name, middle_name))
        if role_name not in ROLE_MAP or login in logins or fio in user_ids:
            raise ValueError(f"Users, строка {number}: неизвестная роль или неоднозначный пользователь")
        user_id = len(users) + 1
        logins.add(login)
        user_ids[fio] = user_id
        users.append((user_id, last_name, first_name, middle_name, login, ROLE_MAP[role_name][0]))

    stock_items = []
    stock_ids = {}
    for number, row in read_rows(source, "Stock_Items_import.xlsx", 4):
        name, manufacturer, size, quantity = row
        key = model_key(name, manufacturer)
        size = decimal_value(size)
        if key not in product_ids or size not in size_ids:
            raise ValueError(f"Stock_Items, строка {number}: неизвестная модель/размер {key}, {size}")
        stock_key = (product_ids[key], size_ids[size])
        if stock_key in stock_ids:
            raise ValueError(f"Stock_Items, строка {number}: повтор модель + размер")
        stock_id = len(stock_items) + 1
        stock_ids[stock_key] = stock_id
        stock_items.append((stock_id, stock_key[0], stock_key[1], integer_value(quantity)))

    orders = {}
    order_items = []
    item_keys = set()
    for number, row in read_rows(source, "Orders_import.xlsx", 9):
        order_id, order_date, fio, category, name, manufacturer, size, quantity, price = row
        order_id = integer_value(order_id, minimum=1)
        order_date = date_value(order_date)
        fio = normalize(fio)
        key = model_key(name, manufacturer)
        size = decimal_value(size)
        if fio not in user_ids or key not in product_ids or size not in size_ids:
            raise ValueError(f"Orders, строка {number}: неизвестный клиент/модель/размер")
        stock_key = (product_ids[key], size_ids[size])
        if stock_key not in stock_ids:
            raise ValueError(f"Orders, строка {number}: товарная позиция отсутствует в остатках")
        product = products[product_ids[key] - 1]
        if normalize(category) != product[1]:
            raise ValueError(f"Orders, строка {number}: категория не совпадает с моделью")
        header = (order_id, order_date, user_ids[fio])
        if order_id in orders and orders[order_id] != header:
            raise ValueError(f"Orders, строка {number}: противоречивые дата/клиент заказа {order_id}")
        orders[order_id] = header
        item_key = (order_id, stock_ids[stock_key])
        if item_key in item_keys:
            raise ValueError(f"Orders, строка {number}: повторная позиция одного заказа")
        item_keys.add(item_key)
        order_items.append((len(order_items) + 1, order_id, stock_ids[stock_key],
                            integer_value(quantity, minimum=1), decimal_value(price)))

    roles = [(value[0], value[1], name) for name, value in ROLE_MAP.items()]
    tables = [
        ("roles", ("role_id", "role_code", "role_name"), roles),
        ("users", ("user_id", "last_name", "first_name", "middle_name", "login", "role_id"), users),
        ("products", ("product_id", "category", "subcategory", "image_path", "product_name", "manufacturer", "description", "composition", "price"), products),
        ("sizes", ("size_id", "size_value"), sizes),
        ("stock_items", ("stock_item_id", "product_id", "size_id", "quantity_available"), stock_items),
        ("orders", ("order_id", "order_date", "user_id"), list(orders.values())),
        ("order_items", ("order_item_id", "order_id", "stock_item_id", "quantity", "unit_price"), order_items),
    ]
    sql = ["-- Только для пустых таблиц после schema.sql. Исторические заказы НЕ списывают остатки повторно.\n",
           "BEGIN;\n", "SET LOCAL standard_conforming_strings = on;\n"]
    for table, columns, rows in tables:
        if not rows:
            raise ValueError(f"Таблица {table} пуста: проверь исходные файлы")
        sql.append(insert_statement(table, columns, rows))
    for table, columns, rows in tables:
        primary_key = columns[0]
        maximum = max(row[0] for row in rows)
        sql.append(f"SELECT setval(pg_get_serial_sequence('public.{table}', '{primary_key}'), {maximum}, true);\n")
    sql.append("COMMIT;\n")
    counts = {table: len(rows) for table, columns, rows in tables}
    counts["available_pairs"] = sum(row[3] for row in stock_items)
    counts["ordered_pairs"] = sum(row[3] for row in order_items)
    return "\n".join(sql), counts


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True, help="Папка с пятью исходными XLSX, поиск рекурсивный")
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parent / "database" / "seed.sql")
    args = parser.parse_args()
    try:
        if not args.source.is_dir():
            raise ValueError(f"Папка источников не найдена: {args.source}")
        sql, counts = build_seed(args.source)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        with args.output.open("w", encoding="utf-8", newline="\n") as stream:
            stream.write(sql)
    except (ValueError, OSError) as error:
        parser.exit(1, f"Ошибка подготовки импорта: {error}\n")
    print(f"SQL сохранён: {args.output.resolve()}")
    for table, count in counts.items():
        print(f"{table}: {count}")


if __name__ == "__main__":
    main()
